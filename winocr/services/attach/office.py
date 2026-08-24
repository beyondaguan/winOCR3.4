# -*- coding: utf-8 -*-
"""附件解析器：Word（段落+表格）与 Excel（多工作表，单表限 500 行预览）。"""
from __future__ import annotations

from .base import AttachParser
from ...core.types import Attachment

_MAX_ROWS_PER_SHEET = 500


class WordParser(AttachParser):
    name = "word"
    display_name = "Word 文档"
    kind = "doc"
    extensions = {".docx", ".doc"}

    def available(self) -> bool:
        try:
            import docx  # noqa: F401
            return True
        except ImportError:
            return False

    def parse(self, path: str) -> Attachment:
        if path.lower().endswith(".doc"):
            # 老式二进制 .doc，python-docx 读不了，明确告知而不是抛难懂的异常
            return self.make(path, error="不支持旧版 .doc，请另存为 .docx")
        try:
            import docx
        except ImportError:
            return self.make(path, error="未安装 python-docx")
        try:
            d = docx.Document(path)
            lines = [p.text for p in d.paragraphs]
            for table in d.tables:               # 表格按 TSV 展平，模型好读
                for row in table.rows:
                    lines.append("\t".join(c.text for c in row.cells))
            return self.make(path, text="\n".join(lines))
        except Exception as e:
            return self.make(path, error=f"Word 解析失败: {e}")


class ExcelParser(AttachParser):
    name = "excel"
    display_name = "Excel 表格"
    kind = "xlsx"
    extensions = {".xlsx", ".xlsm", ".xls"}

    def available(self) -> bool:
        try:
            import openpyxl  # noqa: F401
            return True
        except ImportError:
            return False

    def parse(self, path: str) -> Attachment:
        if path.lower().endswith(".xls"):
            return self.make(path, error="不支持旧版 .xls，请另存为 .xlsx")
        try:
            import openpyxl
        except ImportError:
            return self.make(path, error="未安装 openpyxl")
        wb = None
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"=== 工作表: {ws.title} "
                           f"({ws.max_row}行 x {ws.max_column}列) ===")
                count = 0
                for row in ws.iter_rows(values_only=True):
                    line = "\t".join("" if c is None else str(c) for c in row).rstrip()
                    if line.strip():
                        out.append(line)
                        count += 1
                    if count >= _MAX_ROWS_PER_SHEET:
                        out.append("...（该表后续行已省略）")
                        break
            return self.make(path, text="\n".join(out))
        except Exception as e:
            return self.make(path, error=f"Excel 解析失败: {e}")
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass
